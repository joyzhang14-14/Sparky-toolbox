"""Office 文件读取（规格 6.5、6.6）。

- DOCX：段落/表格；XLSX：sheet/单元格/公式及已有缓存值（不执行宏、不声称已重算）；PPTX：slide/文本/表格。
- 旧 DOC/XLS/PPT：LibreOffice 无头模式转换成 DOCX/XLSX/PPTX 后按上面读取；有超时、输出尺寸上限、禁用宏与外链更新。
- 嵌入脚本、宏、外链更新、自动执行全部禁用（只读取 XML 内容，不触发任何执行）。
"""
from __future__ import annotations

import asyncio
import os
import shutil
import tempfile
from pathlib import Path

from owa_research.core.document import PageSpan, ParsedDocument, Section, TableBlock
from owa_research.core.errors import ErrorCode, ToolError

_LEGACY = {"doc": "docx", "xls": "xlsx", "ppt": "pptx", "rtf": "docx", "odt": "docx", "ods": "xlsx", "odp": "pptx"}


# ---------- DOCX ----------
def parse_docx(data: bytes) -> ParsedDocument:
    import docx  # python-docx，懒加载
    from docx.document import Document as _Doc
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    import io

    d: _Doc = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    sections: list[Section] = []
    tables: list[TableBlock] = []
    pos = 0
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith("heading"):
                level = int("".join(ch for ch in style if ch.isdigit()) or 1)
                line = "#" * min(level, 6) + " " + text
                sections.append(Section(id=f"s{len(sections)}", title=text, level=level, char_start=pos, char_end=pos))
            else:
                line = text
            chunk = line + "\n\n"
            parts.append(chunk)
            pos += len(chunk)
        elif child.tag == qn("w:tbl"):
            t = Table(child, d)
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                tables.append(TableBlock(index=len(tables), rows=rows, headers=rows[0], method="office"))
                md = _rows_to_md(rows) + "\n\n"
                parts.append(md)
                pos += len(md)
    for a, b in zip(sections, sections[1:]):
        a.char_end = b.char_start
    if sections:
        sections[-1].char_end = pos
    text = "".join(parts)
    if not sections:
        sections = [Section(id="s0", title="(全文)", level=0, char_start=0, char_end=len(text))]
    title = (d.core_properties.title or None) if d.core_properties else None
    return ParsedDocument(kind="docx", title=title, text=text, sections=sections, tables=tables, method="python-docx")


# ---------- XLSX ----------
def parse_xlsx(data: bytes, *, max_rows: int = 2000, max_cols: int = 60) -> ParsedDocument:
    import io

    import openpyxl
    from openpyxl.utils import get_column_letter

    # 两次打开：data_only=True 取缓存值，data_only=False 取公式；keep_vba=False 不加载宏，keep_links=False 不更新外链
    wb_vals = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    wb_form = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    parts: list[str] = []
    sections: list[Section] = []
    tables: list[TableBlock] = []
    warnings: list[str] = []
    pos = 0
    for ws_v, ws_f in zip(wb_vals.worksheets, wb_form.worksheets):
        rows: list[list[str]] = []
        formulas: dict[str, str] = {}
        for r_i, (row_v, row_f) in enumerate(zip(ws_v.iter_rows(values_only=True), ws_f.iter_rows(values_only=True)), start=1):
            if r_i > max_rows:
                warnings.append(f"sheet {ws_v.title} 超过 {max_rows} 行，已截断")
                break
            cells = []
            for c_i, (v, f) in enumerate(zip(row_v[:max_cols], row_f[:max_cols]), start=1):
                cells.append("" if v is None else str(v))
                if isinstance(f, str) and f.startswith("="):
                    formulas[f"{get_column_letter(c_i)}{r_i}"] = f
            if any(cells):
                rows.append(cells)
        header = f"## sheet: {ws_v.title}\n"
        sections.append(Section(id=f"s{len(sections)}", title=ws_v.title, level=2, char_start=pos, char_end=pos))
        md = header + (_rows_to_md(rows) if rows else "(空)") + "\n\n"
        parts.append(md)
        pos += len(md)
        sections[-1].char_end = pos
        if rows:
            tables.append(TableBlock(index=len(tables), rows=rows, headers=rows[0], title=ws_v.title, sheet=ws_v.title, cell_origin="A1", method="office",
                                     footnotes=[f"公式 {k}: {v}" for k, v in list(formulas.items())[:200]]))
    wb_vals.close()
    wb_form.close()
    return ParsedDocument(kind="xlsx", title=None, text="".join(parts), sections=sections, tables=tables, warnings=warnings, method="openpyxl",
                          meta={"note": "数值为文件内缓存值，未重算；公式列在表格 footnotes"})


# ---------- PPTX ----------
def parse_pptx(data: bytes) -> ParsedDocument:
    import io

    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    sections: list[Section] = []
    tables: list[TableBlock] = []
    spans: list[PageSpan] = []
    pos = 0
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = [f"## slide {i}"]
        title = None
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    if title is None and getattr(shape, "is_placeholder", False) and shape.placeholder_format.type is not None and "TITLE" in str(shape.placeholder_format.type):
                        title = t
                    lines.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                if rows:
                    tables.append(TableBlock(index=len(tables), rows=rows, headers=rows[0], page=i, method="office"))
                    lines.append(_rows_to_md(rows))
        chunk = "\n".join(lines) + "\n\n"
        spans.append(PageSpan(page=i, char_start=pos, char_end=pos + len(chunk)))
        sections.append(Section(id=f"s{len(sections)}", title=title or f"slide {i}", level=2, char_start=pos, char_end=pos + len(chunk), page=i))
        parts.append(chunk)
        pos += len(chunk)
    return ParsedDocument(kind="pptx", title=None, text="".join(parts), sections=sections, tables=tables, pages=spans, method="python-pptx")


# ---------- 旧格式转换 ----------
async def convert_legacy(data: bytes, ext: str, *, timeout: float = 120, max_output_bytes: int = 100 * 1024 * 1024) -> tuple[bytes, str]:
    """LibreOffice 无头转换 doc/xls/ppt → docx/xlsx/pptx。返回 (bytes, 新扩展名)。"""
    target = _LEGACY.get(ext.lower())
    if target is None:
        raise ToolError(ErrorCode.UNSUPPORTED_FORMAT, f"不支持的旧格式 .{ext}", stage="parse")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise ToolError(ErrorCode.EXTRACTION_FAILED, "研究容器缺少 LibreOffice，无法转换旧格式", stage="parse")
    with tempfile.TemporaryDirectory(prefix="owa-lo-") as tmp:
        src = Path(tmp) / f"input.{ext.lower()}"
        src.write_bytes(data)
        env = {"HOME": tmp, "PATH": os.environ.get("PATH", ""), "SAL_USE_VCLPLUGIN": "svp"}
        # 独立 profile（-env:UserInstallation）避免锁；--norestore 不恢复；宏与外链：转换时不执行宏（无头、无 --enable-macro），
        # 且 LibreOffice 默认对 Untrusted 文档禁用宏；不传 --update 类参数，外链不刷新。
        cmd = [soffice, "--headless", "--invisible", "--nologo", "--norestore", "--nolockcheck",
               f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", target, "--outdir", tmp, str(src)]
        try:
            proc = await asyncio.create_subprocess_exec(*cmd, env=env, stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.PIPE)
            _, err = await asyncio.wait_for(proc.communicate(), timeout=timeout)
        except asyncio.TimeoutError:
            proc.kill()
            raise ToolError(ErrorCode.EXTRACTION_FAILED, f"LibreOffice 转换超过 {timeout} 秒", stage="parse") from None
        out = Path(tmp) / f"input.{target}"
        if proc.returncode != 0 or not out.exists():
            raise ToolError(ErrorCode.EXTRACTION_FAILED, f"LibreOffice 转换失败（退出码 {proc.returncode}）: {err.decode(errors='replace')[-200:]}", stage="parse")
        if out.stat().st_size > max_output_bytes:
            raise ToolError(ErrorCode.EXTRACTION_FAILED, "转换产物超过尺寸上限", stage="parse")
        return out.read_bytes(), target


def _rows_to_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    esc = lambda s: s.replace("|", "\\|").replace("\n", " ")  # noqa: E731
    lines = ["| " + " | ".join(esc(c) for c in norm[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in norm[1:]]
    return "\n".join(lines)
